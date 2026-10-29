"""Streaming export writers (see :mod:`ovid2.exports`)."""

from __future__ import annotations

import csv
import os
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from ovid2.drivers.base import json_scalar

__all__ = [
    "EXCEL_MAX_ROWS",
    "EXPORT_FORMATS",
    "HEADER_CASES",
    "ExportError",
    "ExportResult",
    "apply_header_case",
    "export_rows",
    "media_type_for",
]

#: Supported export formats → (file extension, media type).
EXPORT_FORMATS: dict[str, tuple[str, str]] = {
    "csv": ("csv", "text/csv"),
    "xlsx": ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf": ("pdf", "application/pdf"),
    "docx": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
}

#: Header-case options.
HEADER_CASES: tuple[str, ...] = ("asis", "upper", "lower")

#: Excel's hard sheet limit minus the header row.
EXCEL_MAX_ROWS = 1_048_575

#: Longest cell text rendered in PDF / Word (documents, not data dumps).
_DOC_CELL_MAX = 200


class ExportError(ValueError):
    """Raised for an invalid export request or a missing optional library."""


@dataclass(frozen=True)
class ExportResult:
    """Outcome of :func:`export_rows`.

    Attributes:
        path: The written file.
        format: The export format.
        rows: Data rows written.
        columns: Column count.
        truncated: ``True`` when a format limit stopped the export early.
    """

    path: Path
    format: str
    rows: int
    columns: int
    truncated: bool = False

    def to_dict(self) -> dict[str, Any]:
        """Return a JSON-serialisable dict."""
        return {
            "path": str(self.path),
            "format": self.format,
            "rows": self.rows,
            "columns": self.columns,
            "truncated": self.truncated,
            "size_bytes": self.path.stat().st_size if self.path.exists() else None,
        }


def media_type_for(fmt: str) -> str:
    """Return the HTTP media type of an export format."""
    return EXPORT_FORMATS.get(fmt, ("bin", "application/octet-stream"))[1]


def apply_header_case(columns: Iterable[str], case: str) -> list[str]:
    """Return ``columns`` transformed by ``case`` (``asis`` / ``upper`` / ``lower``)."""
    cols = [str(c) for c in columns]
    if case == "upper":
        return [c.upper() for c in cols]
    if case == "lower":
        return [c.lower() for c in cols]
    return cols


def _cell(value: Any) -> Any:
    """Normalise one value for writing (JSON-scalar, ``None`` → empty)."""
    v = json_scalar(value)
    return "" if v is None else v


def _validate(fmt: str, delimiter: str, header_case: str) -> None:
    if fmt not in EXPORT_FORMATS:
        raise ExportError(f"Unknown export format {fmt!r}; choose one of {sorted(EXPORT_FORMATS)}.")
    if header_case not in HEADER_CASES:
        raise ExportError(f"Unknown header case {header_case!r}; choose one of {HEADER_CASES}.")
    if fmt == "csv" and (len(delimiter) != 1 or delimiter in ("\r", "\n", '"')):
        raise ExportError("The CSV delimiter must be one character (not a quote or newline).")


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------


def _write_csv(
    out: Path, header: list[str], batches: Iterator[list[tuple[Any, ...]]], delimiter: str
) -> tuple[int, bool]:
    n = 0
    with out.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(header)
        for batch in batches:
            writer.writerows([_cell(v) for v in row] for row in batch)
            n += len(batch)
    return n, False


def _write_xlsx(
    out: Path, header: list[str], batches: Iterator[list[tuple[Any, ...]]], sheet: str
) -> tuple[int, bool]:
    import xlsxwriter

    wb = xlsxwriter.Workbook(str(out), {"constant_memory": True, "strings_to_urls": False})
    try:
        ws = wb.add_worksheet((sheet or "data")[:31])
        bold = wb.add_format({"bold": True})
        ws.write_row(0, 0, header, bold)
        ws.freeze_panes(1, 0)
        n = 0
        truncated = False
        for batch in batches:
            for row in batch:
                if n >= EXCEL_MAX_ROWS:
                    truncated = True
                    break
                ws.write_row(n + 1, 0, [_cell(v) for v in row])
                n += 1
            if truncated:
                break
    finally:
        wb.close()
    return n, truncated


def _doc_text(value: Any) -> str:
    text = str(_cell(value))
    return text if len(text) <= _DOC_CELL_MAX else text[: _DOC_CELL_MAX - 1] + "…"


def _write_pdf(
    out: Path,
    header: list[str],
    batches: Iterator[list[tuple[Any, ...]]],
    title: str,
) -> tuple[int, bool]:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A3, landscape, letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    except ImportError as exc:
        raise ExportError("PDF export needs reportlab (pip install 'ovid2[export]').") from exc

    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=6.5, leading=8)
    head_style = styles["BodyText"].clone(
        "head", fontSize=7, leading=8.5, fontName="Helvetica-Bold"
    )
    page = landscape(A3 if len(header) > 12 else letter)
    doc = SimpleDocTemplate(
        str(out),
        pagesize=page,
        leftMargin=24,
        rightMargin=24,
        topMargin=28,
        bottomMargin=28,
        title=title,
    )
    data: list[list[Any]] = [[Paragraph(_escape(h), head_style) for h in header]]
    n = 0
    for batch in batches:
        for row in batch:
            data.append([Paragraph(_escape(_doc_text(v)), cell_style) for v in row])
            n += 1
    width = page[0] - 48
    col_w = [width / max(1, len(header))] * len(header)
    table = LongTable(data, colWidths=col_w, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8edf3")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c2cc")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f9fb")]),
            ]
        )
    )
    meta = Paragraph(_escape(f"{title} — {n} row(s), {len(header)} column(s)"), styles["Heading4"])
    doc.build([meta, Spacer(1, 6), table])
    return n, False


def _escape(text: str) -> str:
    """Escape text for a reportlab ``Paragraph`` (mini-HTML)."""
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _write_docx(
    out: Path,
    header: list[str],
    batches: Iterator[list[tuple[Any, ...]]],
    title: str,
) -> tuple[int, bool]:
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Pt
    except ImportError as exc:
        raise ExportError("Word export needs python-docx (pip install 'ovid2[export]').") from exc

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    document.add_heading(title, level=2)
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, header, strict=False):
        cell.text = name
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    n = 0
    for batch in batches:
        for row in batch:
            cells = table.add_row().cells
            for cell, value in zip(cells, row, strict=False):
                cell.text = _doc_text(value)
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(7.5)
            n += 1
    document.add_paragraph(f"{n} row(s), {len(header)} column(s).")
    document.save(str(out))
    return n, False


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def export_rows(
    columns: list[str],
    batches: Iterator[list[tuple[Any, ...]]],
    out: Path,
    *,
    fmt: str = "csv",
    delimiter: str = ",",
    header_case: str = "asis",
    title: str = "",
) -> ExportResult:
    """Write ``(columns, batches)`` to ``out`` in ``fmt``.

    The file is written to ``<out>.part`` and renamed on success so a
    reader never sees a half-written export. Mode ``0o600`` (exports may
    hold patient data).

    Raises:
        ExportError: Bad format / delimiter / header case, or a missing
            optional library.
    """
    _validate(fmt, delimiter, header_case)
    header = apply_header_case(columns, header_case)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(out.name + ".part")
    old_umask = os.umask(0o077)
    try:
        if fmt == "csv":
            rows, truncated = _write_csv(tmp, header, batches, delimiter)
        elif fmt == "xlsx":
            rows, truncated = _write_xlsx(tmp, header, batches, title)
        elif fmt == "pdf":
            rows, truncated = _write_pdf(tmp, header, batches, title or "Export")
        else:
            rows, truncated = _write_docx(tmp, header, batches, title or "Export")
        os.chmod(tmp, 0o600)
        os.replace(tmp, out)
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise
    finally:
        os.umask(old_umask)
    return ExportResult(path=out, format=fmt, rows=rows, columns=len(header), truncated=truncated)


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit("Library module; import it.")
